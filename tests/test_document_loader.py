import json
import tempfile
import unittest
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfWriter

from backend.knowledge_base.document_loader import load_document
from backend.knowledge_base.vector_store import VectorStore


class DocumentLoaderTest(unittest.TestCase):
    def test_long_document_is_split_into_indexable_knowledge_chunks(self):
        store = VectorStore()
        documents = [Document(
            page_content=("# 技术架构\n接口约束和依赖说明。\n\n" * 80),
            metadata={"source_file": "架构说明.md", "stored_file": "stored_架构说明.md"},
        )]

        chunks = store.split_documents(documents)

        self.assertGreater(len(chunks), 1)
        self.assertEqual([item.metadata["chunk_index"] for item in chunks], list(range(len(chunks))))
        self.assertTrue(all(item.metadata["source_file"] == "架构说明.md" for item in chunks))
        self.assertTrue(all(item.metadata["chunk_count"] == len(chunks) for item in chunks))

    def test_csv_rows_keep_source_and_row_metadata(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "项目成员.csv"
            path.write_text("姓名,岗位\n张三,开发\n李四,测试\n", encoding="utf-8")

            documents = load_document(str(path))

        self.assertEqual(len(documents), 2)
        self.assertIn("姓名: 张三", documents[0].page_content)
        self.assertEqual(documents[0].metadata["source_file"], "项目成员.csv")
        self.assertEqual(documents[1].metadata["row"], 1)

    def test_docx_and_json_use_standalone_loaders(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            docx_path = Path(temp_dir) / "交接说明.docx"
            json_path = Path(temp_dir) / "图谱.json"
            docx = DocxDocument()
            docx.add_heading("交接说明", level=1)
            docx.add_paragraph("部署前必须完成回归测试。")
            docx.save(docx_path)
            json_path.write_text(json.dumps({"节点": ["部署"]}, ensure_ascii=False), encoding="utf-8")

            docx_documents = load_document(str(docx_path))
            json_documents = load_document(str(json_path))

        self.assertIn("部署前必须完成回归测试", docx_documents[0].page_content)
        self.assertEqual(docx_documents[0].metadata["source_file"], "交接说明.docx")
        self.assertIn('"节点"', json_documents[0].page_content)

    def test_pdf_pages_keep_source_and_page_metadata(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "需求说明.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=300, height=200)
            with path.open("wb") as handle:
                writer.write(handle)

            documents = load_document(str(path))

        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].metadata["source_file"], "需求说明.pdf")
        self.assertEqual(documents[0].metadata["page"], 0)

    def test_staging_file_uses_original_filename_hint_for_loader_and_metadata(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "upload.pending"
            path.write_text("暂存文件应按原始扩展名解析。", encoding="utf-8")

            documents = load_document(str(path), "guide.md")

        self.assertIn("按原始扩展名解析", documents[0].page_content)
        self.assertEqual(documents[0].metadata["source_file"], "guide.md")
        self.assertEqual(documents[0].metadata["source"], "guide.md")


if __name__ == "__main__":
    unittest.main()
